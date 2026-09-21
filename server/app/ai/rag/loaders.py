import io
import logging
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional

logger = logging.getLogger(__name__)

# Minimum average characters per page required to consider native text extraction sufficient
SCANNED_PAGE_CHAR_THRESHOLD = 30


@dataclass
class DocumentItem:
    """Represents an extracted page or section from an in-memory document."""
    content: str
    page_number: int
    metadata: Dict[str, Any] = field(default_factory=dict)


def _try_ocr_on_pdf_page(page, page_num: int) -> str:
    """
    Attempts OCR on embedded images of a scanned PDF page only when necessary.
    Gracefully falls back if pytesseract or system tesseract binaries are not available.
    """
    try:
        import pytesseract
        from PIL import Image

        ocr_texts = []
        # Extract images embedded on this page using pypdf
        images = getattr(page, "images", [])
        for img_obj in images:
            try:
                img_bytes = img_obj.data
                with Image.open(io.BytesIO(img_bytes)) as pil_img:
                    # Convert to RGB if needed (e.g. CMYK, palette)
                    if pil_img.mode not in ("L", "RGB"):
                        pil_img = pil_img.convert("RGB")
                    text = pytesseract.image_to_string(pil_img)
                    if text and text.strip():
                        ocr_texts.append(text.strip())
            except Exception as img_err:
                logger.debug(f"OCR image extraction error on page {page_num}: {img_err}")
                continue

        return "\n".join(ocr_texts).strip()
    except ImportError:
        logger.debug("pytesseract or PIL not installed; skipping OCR fallback.")
        return ""
    except Exception as e:
        logger.warning(f"OCR failed for page {page_num}: {e}")
        return ""


def load_pdf_from_bytes(filename: str, content_bytes: bytes) -> List[DocumentItem]:
    """
    Extracts text from a PDF in-memory.
    First attempts normal native text extraction.
    Only attempts OCR if the PDF appears scanned (insufficient text across pages).
    """
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(content_bytes))
    total_pages = len(reader.pages)
    if total_pages == 0:
        raise ValueError(f"PDF file '{filename}' contains no readable pages.")

    pages_text: List[str] = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
            pages_text.append(text.strip())
        except Exception as e:
            logger.warning(f"Error extracting native text from page {i + 1} of '{filename}': {e}")
            pages_text.append("")

    total_chars = sum(len(t) for t in pages_text)
    avg_chars_per_page = total_chars / total_pages

    # Detect if PDF is scanned (text extraction is insufficient)
    is_scanned = avg_chars_per_page < SCANNED_PAGE_CHAR_THRESHOLD

    items: List[DocumentItem] = []
    for i, page in enumerate(reader.pages):
        page_num = i + 1
        page_text = pages_text[i]

        # Use OCR only when the page text is insufficient and document is scanned
        if is_scanned and len(page_text) < SCANNED_PAGE_CHAR_THRESHOLD:
            ocr_text = _try_ocr_on_pdf_page(page, page_num)
            if ocr_text:
                page_text = f"{page_text}\n{ocr_text}".strip() if page_text else ocr_text

        if not page_text:
            page_text = f"[Empty page or non-extractable content on page {page_num}]"

        items.append(
            DocumentItem(
                content=page_text,
                page_number=page_num,
                metadata={
                    "filename": filename,
                    "file_type": "pdf",
                    "total_pages": total_pages,
                    "is_scanned": is_scanned,
                },
            )
        )

    return items


def load_docx_from_bytes(filename: str, content_bytes: bytes) -> List[DocumentItem]:
    """
    Extracts text from a DOCX document in-memory.
    Extracts paragraphs, headings, bullet points, and tables.
    """
    import docx

    doc = docx.Document(io.BytesIO(content_bytes))
    sections_text: List[str] = []

    current_section: List[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        # If heading, optionally start a new section if current has substantial content
        if paragraph.style and "heading" in paragraph.style.name.lower():
            if current_section and sum(len(t) for t in current_section) > 500:
                sections_text.append("\n\n".join(current_section))
                current_section = []
        current_section.append(text)

    # Also extract text from tables
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate identical adjacent cells (due to merged cells)
                deduped = []
                for c in row_cells:
                    if not deduped or deduped[-1] != c:
                        deduped.append(c)
                table_rows.append(" | ".join(deduped))
        if table_rows:
            table_text = "[Table]\n" + "\n".join(table_rows)
            current_section.append(table_text)

    if current_section:
        sections_text.append("\n\n".join(current_section))

    if not sections_text:
        sections_text = ["[Document contains no readable text paragraphs or tables]"]

    items: List[DocumentItem] = []
    for idx, sec_text in enumerate(sections_text):
        items.append(
            DocumentItem(
                content=sec_text,
                page_number=idx + 1,
                metadata={
                    "filename": filename,
                    "file_type": "docx",
                    "total_sections": len(sections_text),
                },
            )
        )

    return items


def load_document(filename: str, content_bytes: bytes) -> List[DocumentItem]:
    """
    Loads and extracts structured text from PDF or DOCX file content in-memory.
    Never persists the raw file to disk or external storage.
    """
    fname_lower = filename.lower().strip()
    if fname_lower.endswith(".pdf"):
        return load_pdf_from_bytes(filename, content_bytes)
    elif fname_lower.endswith(".docx") or fname_lower.endswith(".doc"):
        return load_docx_from_bytes(filename, content_bytes)
    else:
        raise ValueError(
            f"Unsupported document format for '{filename}'. Only PDF (.pdf) and Word documents (.docx) are supported."
        )
